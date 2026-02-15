"""
 @fileoverview

 Service for exporting reading plans to a single Excel workbook
 with monthly_out and daily_out tabs.

 @author tjl
 @version 1.0.0
 @since February 2026
"""

from pathlib import Path
from typing import Dict, List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from src.models.monthly_reading import MonthlyReading
from src.models.reading_day import ReadingDay
from src.services.reference_formatter import ReferenceFormatter

HEADER_FONT = Font(name='Arial', bold=True, size=11)
HEADER_FILL = PatternFill('solid', fgColor='4472C4')
HEADER_FONT_WHITE = Font(name='Arial', bold=True, size=11, color='FFFFFF')
DATA_FONT = Font(name='Arial', size=10)
THIN_BORDER = Border(
    bottom=Side(style='thin', color='D9D9D9')
)


class PlanExporter:
    @staticmethod
    def to_excel(monthly: Dict[int, List[MonthlyReading]],
                 daily: Dict[int, List[ReadingDay]],
                 output_path: Path) -> None:
        wb = Workbook()
        PlanExporter._write_monthly(wb, monthly)
        PlanExporter._write_daily(wb, daily)
        wb.save(output_path)

    @staticmethod
    def _write_monthly(wb: Workbook, yearly: Dict[int, List[MonthlyReading]]) -> None:
        ws = wb.active
        ws.title = 'monthly_out'
        headers = ['Year', 'Month', 'Grouping', 'Books', 'Chapters', 'Words']
        col_widths = [8, 12, 10, 45, 10, 12]

        for col, (header, width) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = HEADER_FONT_WHITE
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal='center')
            ws.column_dimensions[chr(64 + col)].width = width

        row = 2
        for year in range(1, 5):
            for mr in yearly[year]:
                books = mr.book_names
                ws.cell(row=row, column=1, value=year).font = DATA_FONT
                ws.cell(row=row, column=2, value=mr.month_name).font = DATA_FONT
                ws.cell(row=row, column=3, value=mr.grouping).font = DATA_FONT
                ws.cell(row=row, column=4, value=books).font = DATA_FONT
                ws.cell(row=row, column=5, value=len(mr.chapters)).font = DATA_FONT
                cell = ws.cell(row=row, column=6, value=mr.total_words)
                cell.font = DATA_FONT
                cell.number_format = '#,##0'
                for c in range(1, 7):
                    ws.cell(row=row, column=c).border = THIN_BORDER
                row += 1

        ws.auto_filter.ref = f'A1:F{row - 1}'
        ws.freeze_panes = 'A2'

    @staticmethod
    def _write_daily(wb: Workbook, yearly: Dict[int, List[ReadingDay]]) -> None:
        ws = wb.create_sheet('daily_out')
        headers = ['Year', 'Day', 'References', 'Chapters', 'Words']
        col_widths = [8, 8, 45, 10, 12]

        for col, (header, width) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = HEADER_FONT_WHITE
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal='center')
            ws.column_dimensions[chr(64 + col)].width = width

        row = 2
        for year in range(1, 5):
            for rd in yearly[year]:
                ws.cell(row=row, column=1, value=year).font = DATA_FONT
                ws.cell(row=row, column=2, value=rd.day_number).font = DATA_FONT
                ws.cell(row=row, column=3, value=rd.references).font = DATA_FONT
                ws.cell(row=row, column=4, value=len(rd.chapters)).font = DATA_FONT
                cell = ws.cell(row=row, column=5, value=rd.total_words)
                cell.font = DATA_FONT
                cell.number_format = '#,##0'
                for c in range(1, 6):
                    ws.cell(row=row, column=c).border = THIN_BORDER
                row += 1

        ws.auto_filter.ref = f'A1:E{row - 1}'
        ws.freeze_panes = 'A2'
